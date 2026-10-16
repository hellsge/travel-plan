#!/usr/bin/env python3
"""旅行 YAML 与 xlsx 编辑视图之间的导出、解析和同步。"""

from __future__ import annotations

import argparse
import dataclasses
import hashlib
import os
import queue
import re
import sys
import tempfile
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path
from time import monotonic, sleep
from typing import Any, Callable, Optional
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.writer.excel import ExcelWriter
from pydantic import ValidationError

from trip_model import (
    CANCELLED_STATUS,
    ITEM_TYPES,
    PAID_STATUSES,
    RESERVED_IDS,
    STATUSES,
    TODO_STATUSES,
    TRIP_ID_PATTERN,
    Item,
    Trip,
    TripLoadError,
    _describe,
    compute_stats,
    dump_trip,
    load_trip,
)

ROOT = Path(__file__).resolve().parents[1]

HEADERS = (
    "日期", "开始", "结束", "类型", "项目", "路线 / 备注",
    "单价", "数量", "消费", "开售 / 预约时间", "状态", "地点 / 地图关键词",
)
COLUMN_WIDTHS = {
    "A": 12, "B": 11, "C": 11, "D": 9, "E": 20, "F": 60,
    "G": 10, "H": 10, "I": 10, "J": 17, "K": 10, "L": 30,
}
TITLE_CELL, START_CELL, END_CELL, TRAVELERS_CELL = "A1", "B2", "C2", "E2"
TOTAL_CELL, PAID_CELL, UNPAID_CELL = "G2", "I2", "K2"
NOTES_CELL, TODO_CELL = "B3", "K3"
HEADER_ROW = 4
FIRST_ROW = 5
# 公式、下拉和条件格式覆盖到这一行，给手工追加的事项留出余量。
LAST_EDIT_ROW = 2000
META_SHEET, LISTS_SHEET = "_meta", "_lists"
FIXED_TIMESTAMP = datetime(2026, 1, 1)

# 与 templates/assets/app.css 中 .type-* 和 .trip-table td.status-* 的浅色模式一致。
TYPE_COLORS = {
    "交通": "2673d9", "住宿": "7c4dcc", "游览": "219653", "餐饮": "dc7318",
    "日常": "667085", "购物": "d94f8a", "其他": "80685b",
}
STATUS_COLORS = {
    "无需预订": ("eef2f6", "475467"),
    "待购买": ("fef0b8", "704d00"),
    "待预约": ("fef0b8", "704d00"),
    "已购买": ("d7f5df", "176b32"),
    "已预约": ("d7f5df", "176b32"),
    "已完成": ("d7f5df", "176b32"),
    "已取消": ("e5e7eb", "6b7280"),
}

TIME_FORMAT = "hh:mm"
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
DAY_FORMAT = "[$-804]mm-dd aaa"
SHEET_TITLE_FORBIDDEN = set('[]:*?/\\')


ITEM_FIELDS = ("start", "end", "type", "title", "note", "price", "qty", "cost",
               "booking_at", "status", "place")
ITEM_COLUMNS = {field: index for index, field in enumerate(ITEM_FIELDS, start=2)}
SUMMARY_CELLS = {
    "title": (TITLE_CELL, "标题"),
    "start": (START_CELL, "开始日期"),
    "end": (END_CELL, "结束日期"),
    "travelers": (TRAVELERS_CELL, "人数"),
    "notes": (NOTES_CELL, "备注"),
}
DAY_ERROR = re.compile(r"^days\[(\d+)\]\.date (.*)$")
CLOCK_TEXT = re.compile(r"^(\d{1,2})[:：](\d{2})$")


class FileLockedError(Exception):
    def __init__(self, path: Path):
        hint = "请先关闭 Excel" if path.suffix == ".xlsx" else "文件被占用"
        super().__init__(f"{hint}：{path}")
        self.path = path


class XlsxParseError(Exception):
    def __init__(self, messages: list[str]):
        super().__init__("\n".join(messages))
        self.messages = messages


def digest(trip: Trip) -> str:
    return hashlib.sha256(dump_trip(trip).encode("utf-8")).hexdigest()


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _tint(color: str, ratio: float) -> str:
    """CSS color-mix(color ratio, white) 的等价计算。"""
    channels = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c * ratio + 255 * (1 - ratio)):02x}" for c in channels)


def _sheet_title(title: str) -> str:
    cleaned = "".join(" " if ch in SHEET_TITLE_FORBIDDEN else ch for ch in title).strip()
    cleaned = cleaned[:31].strip()
    return cleaned if cleaned and cleaned not in (META_SHEET, LISTS_SHEET) else "行程"


def _clock(text: str | None) -> time | None:
    if text is None:
        return None
    hour, minute = text.split(":")
    return time(int(hour), int(minute))


def _criteria_sum(statuses: frozenset[str]) -> str:
    keys = ",".join(f'"{s}"' for s in STATUSES if s in statuses)
    return f"SUMPRODUCT(SUMIFS({_col('I')},{_col('K')},{{{keys}}}))"


def _col(letter: str) -> str:
    return f"${letter}${FIRST_ROW}:${letter}${LAST_EDIT_ROW}"


def _write_summary(ws, trip: Trip) -> None:
    bold = Font(bold=True)
    ws.merge_cells("A1:L1")
    ws[TITLE_CELL] = trip.title
    ws[TITLE_CELL].font = Font(bold=True, size=16)

    labels = {"A2": "计划日期", "D2": "人数", "F2": "预计总消费", "H2": "已付",
              "J2": "待付", "A3": "备注", "J3": "待办数"}
    for cell, text in labels.items():
        ws[cell] = text
        ws[cell].font = bold
        ws[cell].alignment = Alignment(horizontal="right")
    ws[START_CELL], ws[END_CELL] = trip.start, trip.end
    ws[START_CELL].number_format = ws[END_CELL].number_format = DATE_FORMAT
    ws[TRAVELERS_CELL] = trip.travelers
    ws.merge_cells(f"{NOTES_CELL}:I3")
    ws[NOTES_CELL] = trip.notes
    ws[NOTES_CELL].alignment = Alignment(wrap_text=True, vertical="top")

    ws[TOTAL_CELL] = f'=SUMIFS({_col("I")},{_col("K")},"<>{CANCELLED_STATUS}")'
    ws[PAID_CELL] = f"={_criteria_sum(PAID_STATUSES)}"
    ws[UNPAID_CELL] = f"={_criteria_sum(TODO_STATUSES)}"
    todo_keys = ",".join(f'"{s}"' for s in STATUSES if s in TODO_STATUSES)
    ws[TODO_CELL] = f"=SUMPRODUCT(COUNTIFS({_col('K')},{{{todo_keys}}}))"
    for cell in (TOTAL_CELL, PAID_CELL, UNPAID_CELL, TODO_CELL):
        ws[cell].font = bold

    for index, text in enumerate(HEADERS, start=1):
        cell = ws.cell(HEADER_ROW, index, text)
        cell.font = Font(bold=True, color="ffffff")
        cell.fill = _fill("344054")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_separator(ws, row: int, day) -> None:
    fill = _fill("dbe7f5")
    for col in range(1, len(HEADERS) + 1):
        ws.cell(row, col).fill = fill
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(HEADERS))
    date_cell = ws.cell(row, 1, day.date)
    date_cell.number_format = DAY_FORMAT
    date_cell.font = Font(bold=True, color="1f3a68")
    label = ws.cell(row, 2, day.label)
    label.font = Font(bold=True, color="1f3a68")


def _write_item(ws, row: int, day_date, item: Item) -> None:
    date_cell = ws.cell(row, 1, day_date)
    date_cell.number_format = "mm-dd"
    date_cell.font = Font(color="98a2b3")
    for col, value in ((2, _clock(item.start)), (3, _clock(item.end))):
        ws.cell(row, col, value).number_format = TIME_FORMAT
    ws.cell(row, 4, item.type)
    ws.cell(row, 5, item.title)
    ws.cell(row, 6, item.note).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row, 7, item.price)
    ws.cell(row, 8, item.qty)
    if item.cost is not None:
        ws.cell(row, 9, item.cost)
    else:
        ws.cell(row, 9, f'=IF(OR(G{row}="",H{row}=""),"",G{row}*H{row})')
    if item.booking_at is not None:
        fmt = DATETIME_FORMAT if isinstance(item.booking_at, datetime) else DATE_FORMAT
        ws.cell(row, 10, item.booking_at).number_format = fmt
    ws.cell(row, 11, item.status)
    ws.cell(row, 12, item.place)


def _write_lists(wb: Workbook) -> None:
    ws = wb.create_sheet(LISTS_SHEET)
    for index in range(24 * 6):
        ws.cell(index + 1, 1, time(index // 6, index % 6 * 10)).number_format = TIME_FORMAT
    for index, kind in enumerate(ITEM_TYPES, start=1):
        ws.cell(index, 2, kind)
    for index, status in enumerate(STATUSES, start=1):
        ws.cell(index, 3, status)
    ws.sheet_state = "hidden"


def _add_validations(ws) -> None:
    # 时间下拉只是候选，允许手输 19:28 这类非整 10 分钟的值。
    times = DataValidation(type="list", formula1=f"{LISTS_SHEET}!$A$1:$A$144",
                           showErrorMessage=False)
    types = DataValidation(type="list", formula1=f"{LISTS_SHEET}!$B$1:$B${len(ITEM_TYPES)}",
                           allow_blank=True)
    statuses = DataValidation(type="list", formula1=f"{LISTS_SHEET}!$C$1:$C${len(STATUSES)}",
                              allow_blank=True)
    for dv, first, last in ((times, "B", "C"), (types, "D", "D"), (statuses, "K", "K")):
        dv.add(f"{first}{FIRST_ROW}:{last}{LAST_EDIT_ROW}")
        ws.add_data_validation(dv)


def _add_conditional_formats(ws) -> None:
    for kind in ITEM_TYPES:
        color = TYPE_COLORS[kind]
        rule = CellIsRule(operator="equal", formula=[f'"{kind}"'],
                          fill=_fill(_tint(color, 0.14)), font=Font(bold=True, color=color))
        ws.conditional_formatting.add(f"D{FIRST_ROW}:D{LAST_EDIT_ROW}", rule)
    for status in STATUSES:
        background, color = STATUS_COLORS[status]
        rule = CellIsRule(operator="equal", formula=[f'"{status}"'],
                          fill=_fill(background), font=Font(bold=True, color=color))
        ws.conditional_formatting.add(f"K{FIRST_ROW}:K{LAST_EDIT_ROW}", rule)


def build_workbook(trip: Trip, trip_digest: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(trip.title)
    _write_summary(ws, trip)

    row = FIRST_ROW
    for day in trip.days:
        _write_separator(ws, row, day)
        row += 1
        for item in day.items:
            _write_item(ws, row, day.date, item)
            row += 1

    _add_validations(ws)
    _add_conditional_formats(ws)
    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = f"A{FIRST_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.print_area = f"A1:L{max(row - 1, FIRST_ROW)}"

    _write_lists(wb)
    meta = wb.create_sheet(META_SHEET)
    meta["A1"], meta["B1"] = "digest", trip_digest
    meta.sheet_state = "hidden"

    wb.properties.created = wb.properties.modified = FIXED_TIMESTAMP
    return wb


def _replace_atomically(path: Path, write: Callable[[Path], None]) -> None:
    """先写同目录临时文件再替换，避免被占用或中断时留下半个文件。"""
    fd, tmp_name = tempfile.mkstemp(dir=path.parent, prefix=f".{path.stem}.", suffix=".tmp")
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        write(tmp)
        os.replace(tmp, path)
    except PermissionError:
        raise FileLockedError(path) from None
    finally:
        tmp.unlink(missing_ok=True)


def write_workbook(wb: Workbook, path: Path) -> None:
    def write(tmp: Path) -> None:
        # wb.save() 会把 modified 改成当前时间，这里直接调用 ExcelWriter 保持固定值。
        with ZipFile(tmp, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            ExcelWriter(wb, archive).write_data()

    _replace_atomically(path, write)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any) -> Any:
    if isinstance(value, str):
        try:
            value = float(value.strip().replace(",", ""))
        except ValueError:
            raise ValueError(f"应为数字，实际为 {value!r}") from None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"应为数字，实际为 {value!r}")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _clock_text(value: Any) -> str:
    """接受 Excel 时间值、日期时间、浮点序列值和 HH:MM 文本，四舍五入到分钟。"""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        seconds = value.hour * 3600 + value.minute * 60 + value.second
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        seconds = (value % 1) * 86400
    elif isinstance(value, str) and (match := CLOCK_TEXT.match(value.strip())):
        hour, minute = int(match[1]), int(match[2])
        if hour > 23 or minute > 59:
            raise ValueError(f"时间应为 HH:MM，实际为 {value!r}")
        seconds = hour * 3600 + minute * 60
    else:
        raise ValueError(f"时间应为 HH:MM，实际为 {value!r}")
    minutes = round(seconds / 60) % (24 * 60)
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def _as_date(value: Any, epoch: datetime) -> Optional[date]:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _booking(value: Any, epoch: datetime) -> Any:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        value = (value + timedelta(seconds=30)).replace(second=0, microsecond=0)
        return value.date() if value.time() == time(0, 0) else value
    return value


def _text(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _cell_ref(row: int, col: int) -> str:
    return f"第 {row} 行 {get_column_letter(col)} 列 {HEADERS[col - 1]}"


def _summary_ref(field: str) -> str:
    cell, name = SUMMARY_CELLS[field]
    letters = cell.rstrip("0123456789")
    return f"第 {cell[len(letters):]} 行 {letters} 列 {name}"


def _trip_sheet(wb: Workbook):
    for ws in wb.worksheets:
        if ws.title not in (META_SHEET, LISTS_SHEET):
            return ws
    raise XlsxParseError(["没有找到行程 Sheet"])


class _RowMap:
    """记录 days/items 来自哪一行，用于把 pydantic 的 loc 换算回单元格。"""

    def __init__(self) -> None:
        self.days: list[int] = []
        self.items: dict[tuple[int, int], int] = {}

    def locate(self, loc: tuple) -> Optional[tuple[int, int]]:
        if len(loc) < 2 or loc[0] != "days" or not isinstance(loc[1], int):
            return None
        day = loc[1]
        if len(loc) >= 4 and loc[2] == "items" and isinstance(loc[3], int):
            # Item 级校验（cost 与 price/qty 冲突）没有字段名，定位到消费列。
            field = loc[4] if len(loc) >= 5 else "cost"
            return self.items[(day, loc[3])], ITEM_COLUMNS.get(field, ITEM_COLUMNS["cost"])
        return self.days[day], 2 if len(loc) >= 3 and loc[2] == "label" else 1


def _validation_messages(error: ValidationError, rows: _RowMap) -> list[str]:
    messages: list[str] = []
    for item in error.errors():
        loc, text = item["loc"], _describe(item)
        if not loc and (match := DAY_ERROR.match(text)):
            messages.append(f"第 {rows.days[int(match[1])]} 行 A 列 日期：{match[2]}")
        elif not loc and text.startswith("end "):
            messages.append(f"{_summary_ref('end')}：{text}")
        elif loc and loc[0] in SUMMARY_CELLS:
            messages.append(f"{_summary_ref(loc[0])}：{text}")
        elif loc == ("days",):
            messages.append("至少需要一条日期分隔行（A 列为日期，D–L 列为空）")
        elif (where := rows.locate(loc)) is not None:
            messages.append(f"{_cell_ref(*where)}：{text}")
        else:
            messages.append(text)
    # int | float 这类 Union 字段的同一错误会按分支各报一次。
    return list(dict.fromkeys(messages))


def _parse_item(row: int, cells: dict[str, Any], epoch: datetime) -> tuple[dict, list[str]]:
    item: dict[str, Any] = {}
    errors: list[str] = []

    def convert(field: str, func) -> None:
        value = cells[field]
        if _is_blank(value):
            return
        try:
            item[field] = func(value)
        except ValueError as error:
            errors.append(f"{_cell_ref(row, ITEM_COLUMNS[field])}：{error}")

    convert("start", _clock_text)
    convert("end", _clock_text)
    for field in ("type", "title", "note", "status", "place"):
        convert(field, _text)
    cost = cells["cost"]
    cost_is_formula = isinstance(cost, str) and cost.startswith("=")
    if _is_blank(cells["price"]) and _is_blank(cells["qty"]):
        if not cost_is_formula:
            convert("cost", _number)
    else:
        convert("price", _number)
        convert("qty", _number)
        # 腾讯文档会把 I 列公式转成数值，等于 G×H 时可以安全忽略；否则可能是手填总价，不能丢。
        if not (_is_blank(cost) or cost_is_formula):
            price, qty = item.get("price"), item.get("qty")
            try:
                matches = price is not None and qty is not None and \
                    abs(_number(cost) - price * qty) < 0.005
            except ValueError:
                matches = False
            if not matches:
                errors.append(f"{_cell_ref(row, ITEM_COLUMNS['cost'])}：{cost!r} 与单价×数量"
                              "不一致；只有总价时请清空单价和数量，否则请清空消费列")
    convert("booking_at", lambda value: _booking(value, epoch))
    return item, errors


def _parse_summary(ws, epoch: datetime) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for field, (cell, _name) in SUMMARY_CELLS.items():
        value = ws[cell].value
        if _is_blank(value):
            continue
        if field in ("start", "end"):
            value = _as_date(value, epoch) or value
        elif field == "travelers" and isinstance(value, float) and value.is_integer():
            value = int(value)
        elif field in ("title", "notes"):
            value = _text(value)
        data[field] = value
    return data


def _open_workbook(path: Path) -> Workbook:
    try:
        return load_workbook(path)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise XlsxParseError([f"无法读取 xlsx：{error}"]) from None


def parse_workbook(path: Path, warnings: Optional[list[str]] = None) -> Trip:
    """解析 xlsx；所有错误汇总为 XlsxParseError，warnings 收集不影响结果的提示。"""
    return _parse_loaded(_open_workbook(path), warnings)


def _parse_loaded(wb: Workbook, warnings: Optional[list[str]]) -> Trip:
    ws = _trip_sheet(wb)
    header = tuple(str(ws.cell(HEADER_ROW, col).value or "").strip()
                   for col in range(1, len(HEADERS) + 1))
    if header != HEADERS:
        raise XlsxParseError([f"第 {HEADER_ROW} 行表头应为：{' / '.join(HEADERS)}"])

    data = _parse_summary(ws, wb.epoch)
    errors: list[str] = []
    notes: list[str] = []
    rows = _RowMap()
    days: list[dict[str, Any]] = []
    for row in range(FIRST_ROW, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, len(HEADERS) + 1)]
        row_date = _as_date(values[0], wb.epoch)
        # 没有 label 的分隔行 B–L 也全空，所以要先于空行判断。
        if row_date is not None and all(_is_blank(v) for v in values[3:]):
            day: dict[str, Any] = {"date": row_date, "items": []}
            if not _is_blank(values[1]):
                day["label"] = _text(values[1])
            days.append(day)
            rows.days.append(row)
            continue
        if all(_is_blank(v) for v in values[1:]):
            continue
        if not days:
            errors.append(f"第 {row} 行：事项必须位于某条日期分隔行下方")
            continue
        current = days[-1]
        if row_date is not None and row_date != current["date"]:
            notes.append(f"第 {row} 行 A 列 日期 {row_date} 与所属日期 {current['date']} "
                         "不一致，以上方分隔行为准")
        item, item_errors = _parse_item(row, dict(zip(ITEM_FIELDS, values[1:])), wb.epoch)
        errors.extend(item_errors)
        rows.items[(len(days) - 1, len(current["items"]))] = row
        current["items"].append(item)
    data["days"] = days

    try:
        trip = Trip.model_validate(data)
    except ValidationError as error:
        errors.extend(_validation_messages(error, rows))
    if errors:
        raise XlsxParseError(errors)
    if warnings is not None:
        warnings.extend(notes)
    return trip


NOOP = "已同步"
REEXPORT = "重新导出 xlsx"
TO_YAML = "xlsx → yaml"
TO_XLSX = "yaml → xlsx"
CONFLICT = "冲突"
ERROR = "失败"


@dataclasses.dataclass
class SyncResult:
    trip_id: str
    action: str
    changes: str = ""
    written: list[Path] = dataclasses.field(default_factory=list)
    error: Optional[str] = None
    warnings: list[str] = dataclasses.field(default_factory=list)
    locked: bool = False


def _id_error(trip_id: str) -> Optional[str]:
    if not TRIP_ID_PATTERN.match(trip_id):
        return "文件名只能由小写字母、数字和 - 组成"
    if trip_id in RESERVED_IDS:
        return f"{trip_id} 是站点保留目录名，请换一个文件名"
    return None


def _load_xlsx(path: Path) -> tuple[Trip, Optional[str], list[str]]:
    """返回解析结果、_meta 中记录的 digest 和解析提示。"""
    wb = _open_workbook(path)
    warnings: list[str] = []
    trip = _parse_loaded(wb, warnings)
    meta = wb[META_SHEET]["B1"].value if META_SHEET in wb.sheetnames else None
    return trip, meta if isinstance(meta, str) else None, warnings


def _write_yaml(trip: Trip, path: Path) -> None:
    # 沿用已有文件的换行符，避免 YAML 每次写回都整文件 diff。
    if path.exists():
        newline = "\r\n" if b"\r\n" in path.read_bytes() else "\n"
    else:
        newline = os.linesep
    data = dump_trip(trip).replace("\n", newline).encode("utf-8")
    _replace_atomically(path, lambda tmp: tmp.write_bytes(data))


def _item_changes(old: Trip, new: Trip) -> tuple[int, int, int]:
    """返回 (新增, 删除, 修改)：同一天内容完全相同的不计，其余按标题配对算作修改。"""
    def keyed(trip: Trip) -> Counter:
        return Counter((item.title, day.date, item.model_dump_json())
                       for day in trip.days for item in day.items)

    before, after = keyed(old), keyed(new)
    gone = Counter(key[0] for key in (before - after).elements())
    came = Counter(key[0] for key in (after - before).elements())
    modified = sum((gone & came).values())
    return sum(came.values()) - modified, sum(gone.values()) - modified, modified


def _money(value: float) -> str:
    return f"¥{value:,.2f}"


def _change_summary(old: Optional[Trip], new: Trip) -> str:
    after = compute_stats(new)
    if old is None:
        count = sum(len(day.items) for day in new.days)
        return f"新建 {count} 项事项，总额 {_money(after.total)}，待办 {after.todo_count}"
    before = compute_stats(old)
    added, removed, modified = _item_changes(old, new)
    items = (f"新增 {added} 项、删除 {removed} 项、修改 {modified} 项"
             if added or removed or modified else "事项未变")
    return (f"{items}；总额 {_money(before.total)} → {_money(after.total)}，"
            f"待办 {before.todo_count} → {after.todo_count}")


def _decide(yaml_digest: Optional[str], xlsx_digest: Optional[str], meta: Optional[str],
            prefer: Optional[str]) -> str:
    if xlsx_digest is None:
        return TO_XLSX
    if yaml_digest is None:
        return TO_YAML
    if xlsx_digest == yaml_digest:
        return NOOP if meta == xlsx_digest else REEXPORT
    if prefer is not None:
        return TO_XLSX if prefer == "yaml" else TO_YAML
    if meta == yaml_digest:
        return TO_YAML
    if meta == xlsx_digest:
        return TO_XLSX
    return CONFLICT


def sync_trip(trip_id: str, trips_dir: Path, prefer: Optional[str] = None,
              check: bool = False) -> SyncResult:
    """按 YAML、xlsx 和 _meta 三方 digest 同步一次旅行；check 时只判断不写文件。

    prefer 指定以哪一边为准，另一边无法读取时也会被覆盖。
    """
    if reason := _id_error(trip_id):
        return SyncResult(trip_id, ERROR, error=f"{trip_id}: {reason}")
    yaml_file, xlsx_file = trips_dir / f"{trip_id}.yaml", trips_dir / f"{trip_id}.xlsx"
    if not yaml_file.exists() and not xlsx_file.exists():
        return SyncResult(trip_id, ERROR, error=f"找不到 {yaml_file.name} 或 {xlsx_file.name}")

    errors: list[str] = []
    yaml_trip = xlsx_trip = meta = None
    warnings: list[str] = []
    if yaml_file.exists():
        try:
            yaml_trip = load_trip(yaml_file).trip
        except TripLoadError as error:
            if prefer != "xlsx":
                errors.extend(error.messages)
    if xlsx_file.exists():
        try:
            xlsx_trip, meta, warnings = _load_xlsx(xlsx_file)
        except XlsxParseError as error:
            if prefer != "yaml":
                errors.extend(f"{xlsx_file.name}: {m}" for m in error.messages)
    if not errors and prefer is not None:
        preferred, path = (yaml_trip, yaml_file) if prefer == "yaml" else (xlsx_trip, xlsx_file)
        if preferred is None:
            errors.append(f"找不到 {path.name}")
    if errors:
        return SyncResult(trip_id, ERROR, error="\n".join(errors))

    yaml_digest = digest(yaml_trip) if yaml_trip is not None else None
    xlsx_digest = digest(xlsx_trip) if xlsx_trip is not None else None
    result = SyncResult(trip_id, _decide(yaml_digest, xlsx_digest, meta, prefer),
                        warnings=warnings)
    if result.action == CONFLICT:
        result.changes = f"yaml 相对 xlsx {_change_summary(xlsx_trip, yaml_trip)}"
        result.error = (f"{yaml_file.name} 和 {xlsx_file.name} 都有修改，请运行 sync {trip_id} "
                        "--prefer yaml 或 --prefer xlsx 指定以哪边为准")
        return result
    if result.action == TO_YAML:
        result.changes = _change_summary(yaml_trip, xlsx_trip)
    elif result.action == TO_XLSX:
        result.changes = _change_summary(xlsx_trip, yaml_trip)
    elif result.action == REEXPORT:
        result.changes = "内容未变，刷新同步标记"
    if check or result.action == NOOP:
        return result

    source = xlsx_trip if result.action == TO_YAML else yaml_trip
    try:
        if result.action == TO_YAML:
            _write_yaml(source, yaml_file)
            result.written.append(yaml_file)
        write_workbook(build_workbook(source, digest(source)), xlsx_file)
        result.written.append(xlsx_file)
    except FileLockedError as error:
        # YAML 已写入而 xlsx 被占用时，下次同步会判定为内容一致并补导出 xlsx。
        result.error = str(error)
        result.locked = True
    except OSError as error:
        result.error = str(error)
    return result


def trip_id_for(path: Path) -> Optional[str]:
    """旅行文件对应的 id；Excel 锁文件（~$）和写入中的临时文件（. 开头）返回 None。"""
    if path.suffix not in (".yaml", ".xlsx") or path.name.startswith(("~$", ".")):
        return None
    return path.stem


def sync_all(trips_dir: Path, check: bool = False) -> list[SyncResult]:
    results: list[SyncResult] = []
    trip_ids: set[str] = set()
    for path in sorted(trips_dir.iterdir()):
        if trip_id_for(path) is None or not path.is_file():
            continue
        if reason := _id_error(path.stem):
            results.append(SyncResult(path.stem, ERROR, error=f"{path.name}: {reason}"))
        else:
            trip_ids.add(path.stem)
    results.extend(sync_trip(trip_id, trips_dir, check=check) for trip_id in sorted(trip_ids))
    return results


def _report(result: SyncResult, check: bool, prefix: str = "", out=None) -> bool:
    """打印一次旅行的结果，返回是否算作失败；out 为 None 时提示输出到 stdout。"""
    out = out or sys.stdout
    head = f"{prefix}{result.trip_id}"
    for warning in result.warnings:
        print(f"{head}: 提示：{warning}", file=out)
    detail = f"，{result.changes}" if result.changes else ""
    if result.error:
        print(f"{head}: {result.action}{detail}", file=sys.stderr)
        for line in result.error.splitlines():
            print(f"  {line}", file=sys.stderr)
        return True
    if check and result.action in (TO_YAML, TO_XLSX):
        print(f"{head}: 不一致，需要 {result.action}{detail}", file=sys.stderr)
        return True
    if not check and result.action != NOOP:
        print(f"{head}: {result.action}{detail}", file=out)
    return False


QUIET_SECONDS = 1.5
RETRY_SECONDS = 5.0
POLL_SECONDS = 0.2


class Debouncer:
    """按 trip_id 合并文件事件：最后一次事件后静默满 quiet 秒才到期。"""

    def __init__(self, quiet: float = QUIET_SECONDS, retry: float = RETRY_SECONDS) -> None:
        self.quiet, self.retry = quiet, retry
        self._deadlines: dict[str, float] = {}

    def notify(self, trip_id: str, now: float) -> None:
        self._deadlines[trip_id] = now + self.quiet

    def retry_later(self, trip_id: str, now: float) -> None:
        self._deadlines[trip_id] = now + self.retry

    def due(self, now: float) -> list[str]:
        ready = sorted(trip_id for trip_id, deadline in self._deadlines.items()
                       if deadline <= now)
        for trip_id in ready:
            del self._deadlines[trip_id]
        return ready


class WatchLoop:
    """watch 模式的调度：on_path 可在 watchdog 线程调用，start/tick 在主循环调用。"""

    def __init__(self, trips_dir: Path,
                 sync: Callable[[str, Path], SyncResult] = sync_trip) -> None:
        self.trips_dir = trips_dir
        self._sync = sync
        self._debouncer = Debouncer()
        self._events: queue.SimpleQueue[str] = queue.SimpleQueue()
        self._stamps: dict[str, tuple] = {}
        self._locked: dict[str, str] = {}

    def on_path(self, path: Path) -> None:
        trip_id = trip_id_for(path)
        if trip_id is not None:
            self._events.put(trip_id)

    def start(self, now: float) -> None:
        ids = {trip_id_for(path) for path in self.trips_dir.iterdir() if path.is_file()}
        for trip_id in sorted(ids - {None}):
            self._run(trip_id, now)

    def tick(self, now: float) -> None:
        while True:
            try:
                self._debouncer.notify(self._events.get_nowait(), now)
            except queue.Empty:
                break
        for trip_id in self._debouncer.due(now):
            # Windows 上读文件也可能产生事件，两个文件都没变时不必再解析一遍。
            if self._stamp(trip_id) != self._stamps.get(trip_id):
                self._run(trip_id, now)

    def _stamp(self, trip_id: str) -> tuple:
        def stat(path: Path) -> Optional[tuple[int, int]]:
            try:
                info = path.stat()
            except FileNotFoundError:
                return None
            return info.st_mtime_ns, info.st_size

        return tuple(stat(self.trips_dir / f"{trip_id}{suffix}") for suffix in (".yaml", ".xlsx"))

    def _run(self, trip_id: str, now: float) -> None:
        # 先取文件状态再同步：同步期间用户又保存时，状态对不上，下次事件仍会处理。
        stamp = self._stamp(trip_id)
        self._stamps.pop(trip_id, None)
        if stamp == (None, None):
            return
        prefix = datetime.now().strftime("%H:%M:%S ")
        try:
            result = self._sync(trip_id, self.trips_dir)
        except Exception as error:  # watch 常驻运行，单次失败只打印，不退出。
            print(f"{prefix}{trip_id}: 同步出错：{error!r}", file=sys.stderr)
            return
        if result.locked:
            self._debouncer.retry_later(trip_id, now)
            # 重试期间同一条占用提示只打印一次。
            if self._locked.get(trip_id) == result.error:
                return
            self._locked[trip_id] = result.error
        else:
            self._locked.pop(trip_id, None)
            # 写过文件时不记录状态，让自身写入的事件再同步一次，按 digest 判定为已同步。
            if not result.written:
                self._stamps[trip_id] = stamp
        _report(result, check=False, prefix=prefix)


def run_watch(trips_dir: Path) -> int:
    try:
        from watchdog.events import FileSystemEventHandler
        from watchdog.observers import Observer
    except ImportError:
        print("错误：watch 模式需要 watchdog，请先运行 pip install -r requirements.txt",
              file=sys.stderr)
        return 1
    if not trips_dir.is_dir():
        print(f"错误：找不到目录 {trips_dir}", file=sys.stderr)
        return 1

    loop = WatchLoop(trips_dir)

    class Handler(FileSystemEventHandler):
        def on_any_event(self, event) -> None:
            if event.is_directory:
                return
            for path in (event.src_path, getattr(event, "dest_path", "")):
                if path:
                    loop.on_path(Path(os.fsdecode(path)))

    observer = Observer()
    observer.schedule(Handler(), str(trips_dir), recursive=False)
    observer.start()
    print(f"正在监听 {trips_dir}，按 Ctrl+C 退出")
    try:
        loop.start(monotonic())
        while True:
            sleep(POLL_SECONDS)
            loop.tick(monotonic())
    except KeyboardInterrupt:
        pass
    finally:
        observer.stop()
        observer.join()
    return 0


def parse_args(argv: Optional[list[str]]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="同步旅行 YAML 与 xlsx 编辑视图")
    commands = parser.add_subparsers(dest="command", required=True)
    sync = commands.add_parser("sync", help="按三方 digest 判断方向并同步")
    sync.add_argument("id", nargs="?", help="只同步这一次旅行")
    sync.add_argument("--trips", default=str(ROOT / "trips"), help="旅行文件目录")
    sync.add_argument("--check", action="store_true", help="只检查两边是否一致，不写文件")
    sync.add_argument("--prefer", choices=("yaml", "xlsx"), help="冲突时以哪一边为准，需指定 id")
    sync.add_argument("--watch", action="store_true", help="常驻监听，文件保存后自动同步")
    sync.add_argument("--print-written", action="store_true",
                      help="stdout 只输出写过的文件路径（每行一个），其余提示改到 stderr")
    args = parser.parse_args(argv)
    if args.watch and (args.id or args.check or args.prefer or args.print_written):
        sync.error("--watch 不能与 id、--check、--prefer 或 --print-written 同时使用")
    if args.prefer and not args.id:
        sync.error("--prefer 需要同时指定旅行 id")
    return args


def main(argv: Optional[list[str]] = None) -> int:
    args = parse_args(argv)
    trips_dir = Path(args.trips)
    if args.watch:
        return run_watch(trips_dir)
    if args.id:
        results = [sync_trip(args.id, trips_dir, args.prefer, args.check)]
    else:
        results = sync_all(trips_dir, args.check)
    if not results:
        print(f"错误：{trips_dir} 中没有旅行文件", file=sys.stderr)
        return 1
    info = sys.stderr if args.print_written else sys.stdout
    failed = sum(_report(result, args.check, out=info) for result in results)
    if args.print_written:
        # 失败时也要列出：部分旅行已成功写入，调用方需要知道。
        for result in results:
            for path in result.written:
                print(Path(os.path.relpath(path)).as_posix())
    if failed:
        print(f"{failed} 次旅行{'不一致或' if args.check else '同步'}失败", file=sys.stderr)
        return 1
    print(f"{'检查通过' if args.check else '同步完成'}：{len(results)} 次旅行", file=info)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

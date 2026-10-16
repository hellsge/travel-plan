from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

import trip_model as tm
import trip_xlsx as tx

ROOT = Path(__file__).resolve().parents[1]
VALID_TRIPS, _ = tm.load_trips(ROOT / "trips")


def sep(day, label=None):
    return [datetime(2026, 8, day), label]


def item(day=None, start=None, end=None, kind="游览", title="事项", note=None,
         price=None, qty=None, cost=None, booking=None, status=None, place=None):
    a = datetime(2026, 8, day) if day else None
    return [a, start, end, kind, title, note, price, qty, cost, booking, status, place]


def write_sheet(tmp_path, rows, start=date(2026, 8, 1), end=date(2026, 8, 3)):
    wb = Workbook()
    ws = wb.active
    ws[tx.TITLE_CELL] = "解析测试"
    ws[tx.START_CELL], ws[tx.END_CELL] = start, end
    ws[tx.TRAVELERS_CELL] = 2
    for col, text in enumerate(tx.HEADERS, start=1):
        ws.cell(tx.HEADER_ROW, col, text)
    for offset, values in enumerate(rows):
        for col, value in enumerate(values, start=1):
            ws.cell(tx.FIRST_ROW + offset, col, value)
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


def row_no(index):
    return tx.FIRST_ROW + index


def parse_errors(path):
    with pytest.raises(tx.XlsxParseError) as info:
        tx.parse_workbook(path)
    return info.value.messages


@pytest.mark.parametrize("loaded", VALID_TRIPS, ids=lambda t: t.id)
def test_roundtrip_real_trips(tmp_path, loaded):
    path = tmp_path / f"{loaded.id}.xlsx"
    tx.write_workbook(tx.build_workbook(loaded.trip, tx.digest(loaded.trip)), path)
    warnings = []
    parsed = tx.parse_workbook(path, warnings)
    assert tm.dump_trip(parsed) == tm.dump_trip(loaded.trip)
    assert tx.digest(parsed) == tx.digest(loaded.trip)
    assert warnings == []


def test_basic_fields(tmp_path):
    path = write_sheet(tmp_path, [
        sep(1, "苏州 → 上海"),
        item(1, start="7:28", end="08:05", kind="交通", title=" 高铁 ", price=38, qty=2,
             cost='=IF(OR(G6="",H6=""),"",G6*H6)', booking=datetime(2026, 7, 20, 14, 0),
             status="已购买", place="上海虹桥站"),
        item(1, start=0.5, kind="餐饮", title="午饭", booking=datetime(2026, 7, 31)),
        sep(2),
    ])
    trip = tx.parse_workbook(path)
    assert trip.travelers == 2
    first, second = trip.days[0].items
    assert (first.start, first.end, first.title) == ("07:28", "08:05", "高铁")
    assert (first.price, first.qty, first.cost) == (38, 2, None)
    assert first.booking_at == datetime(2026, 7, 20, 14, 0)
    assert second.start == "12:00"
    assert second.booking_at == date(2026, 7, 31)
    assert second.status == tm.DEFAULT_STATUS
    assert trip.days[0].label == "苏州 → 上海"
    assert trip.days[1].items == []


def test_move_item_to_another_day(tmp_path):
    path = write_sheet(tmp_path, [sep(1), sep(2), item(1, title="挪过来的")])
    warnings = []
    trip = tx.parse_workbook(path, warnings)
    assert trip.days[0].items == []
    assert [i.title for i in trip.days[1].items] == ["挪过来的"]
    assert len(warnings) == 1 and f"第 {row_no(2)} 行 A 列" in warnings[0]


def test_formula_converted_to_value(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1, price=38, qty=2, cost=76)])
    got = tx.parse_workbook(path).days[0].items[0]
    assert (got.price, got.qty, got.cost) == (38, 2, None)


def test_total_only(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1, cost=300)])
    got = tx.parse_workbook(path).days[0].items[0]
    assert (got.price, got.qty, got.cost) == (None, None, 300)


def test_text_numbers_and_times(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1, start="19:28", end="9：05", price="38", qty=" 2 ")])
    got = tx.parse_workbook(path).days[0].items[0]
    assert (got.start, got.end, got.price, got.qty) == ("19:28", "09:05", 38, 2)


def test_blank_rows_skipped(tmp_path):
    path = write_sheet(tmp_path, [
        sep(1), [None] * 12, item(1, title="a"), ["备忘"] + [None] * 11,
        item(1, title="b"), ["  "] * 12,
    ])
    trip = tx.parse_workbook(path)
    assert len(trip.days) == 1
    assert [i.title for i in trip.days[0].items] == ["a", "b"]


def test_item_before_first_separator(tmp_path):
    path = write_sheet(tmp_path, [item(1), sep(1)])
    messages = parse_errors(path)
    assert any(f"第 {row_no(0)} 行" in m for m in messages)


def test_invalid_type_reports_row(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1), item(1, kind="吃饭")])
    messages = parse_errors(path)
    assert len(messages) == 1
    assert messages[0].startswith(f"第 {row_no(2)} 行 D 列 类型：无效值 '吃饭'")


def test_missing_title_reports_row(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1, title=None)])
    assert parse_errors(path) == [f"第 {row_no(1)} 行 E 列 项目：缺少必填字段"]


def test_price_and_total_conflict(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1, price=38, cost=100)])
    messages = parse_errors(path)
    assert len(messages) == 1 and messages[0].startswith(f"第 {row_no(1)} 行 I 列")


def test_bad_number_and_time(tmp_path):
    path = write_sheet(tmp_path, [sep(1), item(1, start="早上", price="三十")])
    messages = parse_errors(path)
    assert any(m.startswith(f"第 {row_no(1)} 行 B 列") for m in messages)
    assert any(m.startswith(f"第 {row_no(1)} 行 G 列") for m in messages)


def test_day_out_of_range(tmp_path):
    path = write_sheet(tmp_path, [sep(1), sep(9)])
    messages = parse_errors(path)
    assert len(messages) == 1
    assert messages[0].startswith(f"第 {row_no(1)} 行 A 列 日期：")
    assert "days[" not in messages[0]


def test_summary_errors(tmp_path):
    path = write_sheet(tmp_path, [sep(1)], start=date(2026, 8, 3), end=date(2026, 8, 1))
    assert any(m.startswith("第 2 行 C 列") for m in parse_errors(path))


def test_no_days(tmp_path):
    path = write_sheet(tmp_path, [])
    assert parse_errors(path)


def test_wrong_header(tmp_path):
    path = write_sheet(tmp_path, [sep(1)])
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb.active.cell(tx.HEADER_ROW, 4, "分类")
    wb.save(path)
    assert any(f"第 {tx.HEADER_ROW} 行" in m for m in parse_errors(path))


def test_not_an_xlsx(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_text("not a zip")
    assert parse_errors(path)

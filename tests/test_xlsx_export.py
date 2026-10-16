from datetime import date, datetime, time

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import trip_model as tm
import trip_xlsx as tx


def make_trip():
    return tm.Trip.model_validate({
        "title": "导出测试",
        "start": date(2026, 8, 1),
        "end": date(2026, 8, 3),
        "travelers": 2,
        "notes": "备注",
        "days": [
            {"date": date(2026, 8, 1), "label": "苏州 → 上海", "items": [
                {"title": "高铁", "type": "交通", "start": "07:28", "end": "07:59",
                 "price": 38, "qty": 2, "booking_at": "2026-07-20 14:00", "status": "已购买"},
                {"title": "酒店", "type": "住宿", "cost": 300, "status": "待购买"},
                {"title": "散步", "type": "游览"},
            ]},
            {"date": date(2026, 8, 2), "items": []},
            {"date": date(2026, 8, 3), "items": [
                {"title": "午饭", "type": "餐饮", "price": 50, "booking_at": "2026-08-01"},
            ]},
        ],
    })


@pytest.fixture
def exported(tmp_path):
    trip = make_trip()
    path = tmp_path / "x.xlsx"
    tx.write_workbook(tx.build_workbook(trip, tx.digest(trip)), path)
    return trip, load_workbook(path)


def rows_of(ws):
    return list(ws.iter_rows(min_row=tx.FIRST_ROW, values_only=True))


def test_digest_is_sha256_of_dump():
    trip = make_trip()
    assert len(tx.digest(trip)) == 64
    assert tx.digest(trip) == tx.digest(tm.Trip.model_validate(trip.model_dump()))


def test_sheets_and_meta(exported):
    trip, wb = exported
    assert wb.sheetnames == ["导出测试", "_lists", "_meta"]
    assert wb["_lists"].sheet_state == "hidden"
    assert wb["_meta"].sheet_state == "hidden"
    assert wb["_meta"]["B1"].value == tx.digest(trip)


def test_summary(exported):
    _, wb = exported
    ws = wb.worksheets[0]
    assert ws[tx.TITLE_CELL].value == "导出测试"
    assert ws[tx.START_CELL].value.date() == date(2026, 8, 1)
    assert ws[tx.END_CELL].value.date() == date(2026, 8, 3)
    assert ws[tx.TRAVELERS_CELL].value == 2
    assert ws[tx.NOTES_CELL].value == "备注"
    for cell in (tx.TOTAL_CELL, tx.PAID_CELL, tx.UNPAID_CELL, tx.TODO_CELL):
        assert str(ws[cell].value).startswith("=")
        assert tm.CANCELLED_STATUS in ws[tx.TOTAL_CELL].value


def test_day_rows_include_empty_day(exported):
    _, wb = exported
    ws = wb.worksheets[0]
    rows = rows_of(ws)
    separators = [(i, r) for i, r in enumerate(rows) if all(v is None for v in r[3:12])]
    assert [r[0].date() for _, r in separators] == [date(2026, 8, 1), date(2026, 8, 2), date(2026, 8, 3)]
    assert separators[0][1][1] == "苏州 → 上海"
    assert separators[1][0] + 1 == separators[2][0]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert f"B{tx.FIRST_ROW}:L{tx.FIRST_ROW}" in merged


def test_item_cells(exported):
    _, wb = exported
    ws = wb.worksheets[0]
    train = tx.FIRST_ROW + 1
    assert ws[f"B{train}"].value == time(7, 28)
    assert ws[f"B{train}"].number_format == "hh:mm"
    assert ws[f"D{train}"].value == "交通"
    assert ws[f"I{train}"].value == f'=IF(OR(G{train}="",H{train}=""),"",G{train}*H{train})'
    assert ws[f"J{train}"].value == datetime(2026, 7, 20, 14, 0)
    assert ws[f"K{train}"].value == "已购买"
    hotel = train + 1
    assert ws[f"G{hotel}"].value is None and ws[f"H{hotel}"].value is None
    assert ws[f"I{hotel}"].value == 300
    walk = hotel + 1
    assert ws[f"I{walk}"].value.startswith("=IF(")
    assert ws[f"K{walk}"].value == tm.DEFAULT_STATUS


def test_data_validations(exported):
    _, wb = exported
    ws = wb.worksheets[0]
    by_col = {}
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            for col in range(rng.min_col, rng.max_col + 1):
                by_col[get_column_letter(col)] = dv
    assert set(by_col) >= {"B", "C", "D", "K"}
    assert by_col["B"] is by_col["C"]
    assert by_col["B"].showErrorMessage is False
    assert by_col["D"].formula1.startswith("_lists!")
    lists = wb["_lists"]
    assert [c.value for c in lists["B"] if c.value] == list(tm.ITEM_TYPES)
    assert [c.value for c in lists["C"] if c.value] == list(tm.STATUSES)
    times = [c.value for c in lists["A"] if c.value is not None]
    assert len(times) == 144 and times[0] == time(0, 0) and times[-1] == time(23, 50)


def test_conditional_formatting(exported):
    _, wb = exported
    ws = wb.worksheets[0]
    rules = {}
    for cf in ws.conditional_formatting:
        col = str(cf.sqref)[0]
        rules[col] = [r.formula[0].strip('"') for r in cf.rules]
    assert sorted(rules["D"]) == sorted(tm.ITEM_TYPES)
    assert sorted(rules["K"]) == sorted(tm.STATUSES)
    assert set(rules) == {"D", "K"}


def test_layout(exported):
    _, wb = exported
    ws = wb.worksheets[0]
    assert ws.freeze_panes == f"A{tx.FIRST_ROW}"
    assert ws.page_setup.orientation == "landscape"
    assert ws.column_dimensions["F"].width > ws.column_dimensions["A"].width


def test_sheet_title_sanitized():
    trip = make_trip().model_copy(update={"title": "a/b:c*d?" + "长" * 40})
    title = tx.build_workbook(trip, "x").worksheets[0].title
    assert len(title) <= 31 and not set(title) & set("[]:*?/\\")


def test_fixed_properties(exported):
    _, wb = exported
    assert wb.properties.created == tx.FIXED_TIMESTAMP
    assert wb.properties.modified == tx.FIXED_TIMESTAMP


def test_write_is_atomic_and_reports_lock(tmp_path, monkeypatch):
    trip = make_trip()
    path = tmp_path / "x.xlsx"
    path.write_bytes(b"old")

    def locked(src, dst):
        raise PermissionError(13, "denied")

    monkeypatch.setattr(tx.os, "replace", locked)
    with pytest.raises(tx.FileLockedError) as info:
        tx.write_workbook(tx.build_workbook(trip, tx.digest(trip)), path)
    assert "请先关闭 Excel" in str(info.value) and "x.xlsx" in str(info.value)
    assert path.read_bytes() == b"old"
    assert list(tmp_path.iterdir()) == [path]
